#!/usr/bin/env python3
"""Convert an Excel workbook into the JSON file used by this static site."""

import argparse
import json
import re
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook


HYPERLINK_RE = re.compile(
    r'^=HYPERLINK\(\s*"((?:[^"]|"")*)"(?:\s*(?:,|;)\s*"(?:[^"]|"")*")?\s*\)$',
    re.IGNORECASE,
)


def serialise(value):
    """Return a JSON-friendly value while keeping spreadsheet values readable."""
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, str):
        match = HYPERLINK_RE.match(value.strip())
        if match:
            # Excel escapes a literal quote inside a formula as two quotes.
            return match.group(1).replace('""', '"')
    return value


def main():
    parser = argparse.ArgumentParser(description="Convert products.xlsx to products.json")
    parser.add_argument("input", nargs="?", default="products.xlsx", help="Excel input file")
    parser.add_argument("-o", "--output", default="products.json", help="JSON output file")
    parser.add_argument("--sheet", help="Worksheet name (defaults to the first sheet)")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        parser.error(f"Input file not found: {input_path}")

    workbook = load_workbook(input_path, data_only=False, read_only=True)
    if args.sheet:
        if args.sheet not in workbook.sheetnames:
            parser.error(f"Sheet not found: {args.sheet}. Available: {', '.join(workbook.sheetnames)}")
        sheet = workbook[args.sheet]
    else:
        sheet = workbook.worksheets[0]

    values = sheet.iter_rows(values_only=True)
    try:
        headers = [serialise(cell) for cell in next(values)]
    except StopIteration:
        headers = []

    # Array-based rows preserve all columns exactly, even if Excel has duplicate or blank headers.
    rows = [[serialise(cell) for cell in row[:len(headers)]] for row in values]
    payload = {"sheet": sheet.title, "headers": headers, "rows": rows}

    output_path = Path(args.output)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(f"Wrote {len(rows)} products from '{sheet.title}' to {output_path}")


if __name__ == "__main__":
    main()
